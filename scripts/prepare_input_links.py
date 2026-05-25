from __future__ import annotations

import argparse
import csv
from pathlib import Path

from openpyxl import load_workbook


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Convert input-link xlsx to CSV.")
    parser.add_argument("--xlsx", default="Input 링크 .xlsx")
    parser.add_argument("--output", default="artifacts/input_links.csv")
    return parser.parse_args()


def main() -> None:
    args = parse_args()
    workbook = load_workbook(args.xlsx, data_only=True)
    worksheet = workbook.active

    rows = list(worksheet.iter_rows(values_only=True))
    if not rows:
        raise RuntimeError(f"No rows found in {args.xlsx}")

    header = [str(value).strip() if value is not None else "" for value in rows[0]]
    try:
        id_col = header.index("번호")
        link_col = header.index("링크")
    except ValueError as exc:
        raise RuntimeError(f"Expected columns ['번호', '링크'], got {header}") from exc

    records: list[dict[str, str]] = []
    for row in rows[1:]:
        raw_video_id = row[id_col]
        raw_url = row[link_col]
        if raw_video_id is None or raw_url is None:
            continue

        video_id = str(raw_video_id).strip()
        if video_id.isdigit():
            video_id = video_id.zfill(4)

        records.append(
            {
                "video_id": video_id,
                "url": str(raw_url).strip(),
            }
        )

    output = Path(args.output)
    output.parent.mkdir(parents=True, exist_ok=True)
    with output.open("w", newline="", encoding="utf-8") as f:
        writer = csv.DictWriter(f, fieldnames=["video_id", "url"])
        writer.writeheader()
        writer.writerows(records)

    print(f"변환 완료: {output}")
    print(f"링크 수: {len(records)}")
    print(f"첫 링크: {records[0] if records else None}")


if __name__ == "__main__":
    main()
